"""Deterministic questionnaire parsing for the Phase 1 input formats."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections.abc import Sequence

from openpyxl import load_workbook

from trustilo.common.schemas import Question, SourceFormat


_QUESTION_HEADERS = frozenset(
    {
        "question",
        "questions",
        "questiontext",
        "questiondescription",
        "controlquestion",
        "query",
        "prompt",
    }
)
_SECTION_HEADERS = frozenset({"section", "sectionname", "category"})


def _normalize_whitespace(value: str) -> str:
    """Collapse whitespace without rewriting words or punctuation."""

    return " ".join(value.split())


def _canonical_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _stable_id(prefix: str, *parts: str | bytes) -> str:
    digest = hashlib.sha256()
    for part in parts:
        encoded = part if isinstance(part, bytes) else part.encode("utf-8")
        digest.update(len(encoded).to_bytes(8, "big"))
        digest.update(encoded)
    return f"{prefix}_{digest.hexdigest()[:24]}"


def _decode_utf8(raw_file: bytes, source_format: SourceFormat) -> str:
    try:
        return raw_file.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(f"{source_format.value.upper()} input must be valid UTF-8 text") from exc


def _column_index(headers: Sequence[object], accepted: frozenset[str], kind: str) -> int | None:
    matches: list[int] = []
    for index, header in enumerate(headers):
        if not isinstance(header, str):
            continue
        if _canonical_header(header.strip()) in accepted:
            matches.append(index)

    if len(matches) > 1:
        raise ValueError(f"Multiple {kind} columns found; the input is ambiguous")
    return matches[0] if matches else None


def _question(
    *,
    tenant_id: str,
    questionnaire_id: str,
    raw_text: str,
    source_format: SourceFormat,
    locator: str,
    source_row_index: int | None = None,
    section: str | None = None,
) -> Question:
    normalized_text = _normalize_whitespace(raw_text)
    if not normalized_text:
        raise ValueError("Question text cannot be blank")

    normalized_section = _normalize_whitespace(section) if section else None
    return Question(
        tenant_id=tenant_id,
        question_id=_stable_id("q", questionnaire_id, locator, normalized_text),
        questionnaire_id=questionnaire_id,
        raw_text=raw_text,
        normalized_text=normalized_text,
        section=normalized_section or None,
        source_format=source_format,
        source_row_index=source_row_index,
    )


def _parse_text(raw_file: bytes, tenant_id: str, questionnaire_id: str) -> list[Question]:
    text = _decode_utf8(raw_file, SourceFormat.TEXT)
    questions: list[Question] = []
    for line_number, raw_text in enumerate(text.splitlines(), start=1):
        if not raw_text.strip():
            continue
        questions.append(
            _question(
                tenant_id=tenant_id,
                questionnaire_id=questionnaire_id,
                raw_text=raw_text,
                source_format=SourceFormat.TEXT,
                locator=f"line:{line_number}",
            )
        )
    return questions


def _parse_csv(raw_file: bytes, tenant_id: str, questionnaire_id: str) -> list[Question]:
    text = _decode_utf8(raw_file, SourceFormat.CSV)
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    try:
        headers = next(reader)
    except StopIteration:
        return []
    except csv.Error as exc:
        raise ValueError(f"Invalid CSV header: {exc}") from exc

    question_index = _column_index(headers, _QUESTION_HEADERS, "question")
    if question_index is None:
        raise ValueError("CSV requires a question column (for example 'Question' or 'Question Text')")
    section_index = _column_index(headers, _SECTION_HEADERS, "section")

    questions: list[Question] = []
    while True:
        source_row_index = reader.line_num + 1
        try:
            row = next(reader)
        except StopIteration:
            break
        except csv.Error as exc:
            raise ValueError(f"Invalid CSV near physical row {source_row_index}: {exc}") from exc

        if not row or all(not value.strip() for value in row):
            continue
        if len(row) != len(headers):
            raise ValueError(
                f"CSV row {source_row_index} has {len(row)} columns; expected {len(headers)}"
            )

        raw_text = row[question_index]
        if not raw_text.strip():
            raise ValueError(f"CSV row {source_row_index} has data but no question text")
        section = row[section_index] if section_index is not None else None
        questions.append(
            _question(
                tenant_id=tenant_id,
                questionnaire_id=questionnaire_id,
                raw_text=raw_text,
                source_format=SourceFormat.CSV,
                locator=f"row:{source_row_index}",
                source_row_index=source_row_index,
                section=section,
            )
        )
    return questions


def _parse_xlsx(raw_file: bytes, tenant_id: str, questionnaire_id: str) -> list[Question]:
    try:
        workbook = load_workbook(io.BytesIO(raw_file), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("Invalid XLSX workbook") from exc

    try:
        worksheet = workbook.active
        if worksheet is None:
            return []

        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            return []

        question_index = _column_index(headers, _QUESTION_HEADERS, "question")
        if question_index is None:
            raise ValueError("XLSX requires a question column (for example 'Question' or 'Question Text')")
        section_index = _column_index(headers, _SECTION_HEADERS, "section")

        questions: list[Question] = []
        for source_row_index, row in enumerate(rows, start=2):
            is_blank = not row or all(
                value is None or (isinstance(value, str) and not value.strip()) for value in row
            )
            if is_blank:
                continue

            raw_text = row[question_index] if question_index < len(row) else None
            if not isinstance(raw_text, str) or not raw_text.strip():
                raise ValueError(f"XLSX row {source_row_index} has data but no text question")

            section_value = (
                row[section_index]
                if section_index is not None and section_index < len(row)
                else None
            )
            if section_value is not None and not isinstance(section_value, str):
                raise ValueError(f"XLSX row {source_row_index} has a non-text section value")

            questions.append(
                _question(
                    tenant_id=tenant_id,
                    questionnaire_id=questionnaire_id,
                    raw_text=raw_text,
                    source_format=SourceFormat.XLSX,
                    locator=f"sheet:{worksheet.title}:row:{source_row_index}",
                    source_row_index=source_row_index,
                    section=section_value,
                )
            )
        return questions
    finally:
        workbook.close()


def parse(raw_file: bytes, source_format: SourceFormat | str, tenant_id: str) -> list[Question]:
    """Parse a questionnaire into ordered canonical ``Question`` records.

    Phase 1 supports UTF-8 plain text, UTF-8 CSV, and the active worksheet
    of an XLSX workbook. PDF parsing is intentionally deferred.
    """

    if not isinstance(raw_file, bytes):
        raise TypeError("raw_file must be bytes")
    if not isinstance(tenant_id, str) or not tenant_id.strip():
        raise ValueError("tenant_id must be a non-empty string")
    try:
        resolved_format = SourceFormat(source_format)
    except ValueError as exc:
        raise ValueError(f"Unsupported source format: {source_format!r}") from exc

    if resolved_format is SourceFormat.PDF:
        raise NotImplementedError("PDF parsing is not implemented in Phase 1")

    questionnaire_id = _stable_id("questionnaire", tenant_id, resolved_format.value, raw_file)
    if resolved_format is SourceFormat.TEXT:
        return _parse_text(raw_file, tenant_id, questionnaire_id)
    if resolved_format is SourceFormat.CSV:
        return _parse_csv(raw_file, tenant_id, questionnaire_id)
    return _parse_xlsx(raw_file, tenant_id, questionnaire_id)
